# -*- coding: utf-8 -*-
"""
多酚数据库双系统后端
- 天然多酚系统 (natural): 7470 个天然多酚化合物
- 油料多酚系统 (oil): 油料籽中检出的多酚 + 文献证据 + UPLC 参数
"""
import os
import io
import sqlite3
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "polyphenol.db")

NATURAL_CLASS_MAP = {
    "Flavonoids": "黄酮类", "黄酮类": "黄酮类", "黄酮木脂素": "黄酮类",
    "Phenolic acids": "酚酸类", "酚酸": "酚酸类",
    "Lignans": "木脂素类", "木脂素": "木脂素类",
    "Stilbenes": "芪类", "芪类": "芪类",
    "Tannins": "单宁类", "单宁": "单宁类",
    "Curcuminoids": "姜黄素类", "姜黄素类": "姜黄素类",
    "Xanthones": "氧杂蒽酮类", "氧杂蒽酮类": "氧杂蒽酮类",
    "Phenolic alcohols": "酚醇类", "酚醇与苯乙醇苷": "酚醇类",
    "Marine phenolics": "海藻多酚", "海藻多酚": "海藻多酚",
    "Olive phenolics": "裂环烯醚萜酚", "橄榄/裂环烯醚萜酚": "裂环烯醚萜酚",
    "酚酰胺": "酚酰胺",
    "多酚醛": "酚醛类",
    "Related phenolics": "其他酚类", "边界相关天然酚类": "其他酚类",
    "FooDB phenolic candidates": "其他酚类", "Phenolic compounds": "其他酚类",
}

OIL_CLASS_MAP = {
    "黄酮": "黄酮类", "黄酮类": "黄酮类", "黄酮醇": "黄酮类", "黄酮苷": "黄酮类",
    "二氢黄酮": "黄酮类", "二氢黄酮醇": "黄酮类", "其它类黄酮": "黄酮类",
    "异黄酮": "黄酮类", "花青素": "黄酮类", "双黄酮": "黄酮类",
    "黄烷醇类": "黄酮类", "橙酮类": "黄酮类", "查耳酮": "黄酮类",
    "查耳酮/喹诺查尔酮": "黄酮类", "原花青素": "黄酮类",
    "儿茶素": "黄酮类", "表儿茶素": "黄酮类",
    "酚酸类": "酚酸类", "酚酸": "酚酸类", "酚酸苷": "酚酸类", "酚酸酯": "酚酸类",
    "木脂素": "木脂素类", "木脂素类": "木脂素类", "木脂酰胺": "木脂素类",
    "香豆素": "香豆素类", "香豆素类": "香豆素类",
    "鞣质": "单宁类", "单宁": "单宁类", "单宁类": "单宁类",
    "茋类": "芪类", "芪类": "芪类", "二苯乙烯类": "芪类",
    "裂环烯醚萜酚": "裂环烯醚萜酚",
    "简单酚": "其他酚类", "萘醌类": "其他酚类", "酚醛类化合物": "其他酚类",
}

def norm_class(raw, mapping):
    if raw is None:
        return "其他酚类"
    key = str(raw).strip()
    return mapping.get(key, "其他酚类")


def init_schema():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS natural_compounds (
        name_en TEXT PRIMARY KEY, name_cn TEXT,
        std_class TEXT, raw_class TEXT, secondary_class TEXT,
        family TEXT, inclusion_level TEXT, cas TEXT, formula TEXT,
        representative_plants TEXT, common_sources TEXT,
        mce TEXT, targetmol TEXT, phenol_explorer TEXT, foodb TEXT,
        mce_url TEXT, targetmol_url TEXT, pe_url TEXT,
        smiles TEXT, smiles_source TEXT, foodb_smiles TEXT,
        foodb_public_id TEXT, foodb_evidence_level TEXT,
        foodb_superklass TEXT, foodb_klass TEXT, foodb_subklass TEXT,
        foodb_description TEXT, pubchem_title TEXT)""")
    c.execute("""CREATE TABLE IF NOT EXISTS oil_compounds (
        name TEXT PRIMARY KEY, std_class TEXT, raw_class TEXT,
        sources TEXT, reference TEXT,
        total_count INTEGER, orig_count INTEGER,
        new_lit_count INTEGER, uplc_count INTEGER,
        online_lit_count INTEGER, online_entity_count INTEGER)""")
    c.execute("""CREATE TABLE IF NOT EXISTS oil_evidence (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        compound_name TEXT, std_class TEXT, raw_class TEXT,
        source TEXT, reference TEXT, data_source TEXT, platform TEXT,
        pdf_file TEXT, doi TEXT, url TEXT, snippet TEXT)""")
    c.execute("""CREATE TABLE IF NOT EXISTS oil_uplc (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        compound_name TEXT, std_class TEXT, original_name TEXT,
        raw_class TEXT, primary_class TEXT, secondary_class TEXT,
        formula TEXT, q1 TEXT, q3 TEXT, mw TEXT, ionization TEXT,
        level TEXT, cas TEXT, cpd_id TEXT, fp TEXT, ep TEXT, kegg_map TEXT)""")
    c.execute("""CREATE TABLE IF NOT EXISTS oil_literature (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        crop TEXT, platform TEXT, query_text TEXT, title TEXT,
        doi TEXT, year TEXT, journal TEXT, url TEXT,
        compounds_identified TEXT, count TEXT)""")
    c.execute("CREATE INDEX IF NOT EXISTS idx_oil_ev ON oil_evidence(compound_name)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_oil_up ON oil_uplc(compound_name)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_nat_class ON natural_compounds(std_class)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_oil_class ON oil_compounds(std_class)")
    conn.commit()
    conn.close()

if not os.path.exists(DB_PATH):
    init_schema()
    print("Empty database created.")

app = FastAPI(title="Polyphenol Dual System")
app.add_middleware(CORSMiddleware, allow_origins=["*"],
                   allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

STATIC_DIR = os.path.join(BASE_DIR, "static")
os.makedirs(STATIC_DIR, exist_ok=True)

@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def _safe(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None

def _int(v):
    try:
        return int(v) if v is not None else 0
    except (ValueError, TypeError):
        return 0


@app.get("/api/natural/stats")
def natural_stats():
    conn = get_db()
    c = conn.cursor()
    total = c.execute("SELECT COUNT(*) FROM natural_compounds").fetchone()[0]
    smi = c.execute("SELECT COUNT(*) FROM natural_compounds WHERE smiles IS NOT NULL OR foodb_smiles IS NOT NULL").fetchone()[0]
    n_class = c.execute("SELECT COUNT(DISTINCT std_class) FROM natural_compounds").fetchone()[0]
    class_dist = {}
    for r in c.execute("SELECT std_class, COUNT(*) n FROM natural_compounds GROUP BY std_class ORDER BY n DESC"):
        class_dist[r[0]] = r[1]
    cross = {}
    for col, key in [("mce","MCE"),("targetmol","TargetMol"),("phenol_explorer","PhenolExplorer"),("foodb","FooDB")]:
        cross[key] = c.execute(f"SELECT COUNT(*) FROM natural_compounds WHERE {col}='Yes'").fetchone()[0]
    conn.close()
    return {
        "total_compounds": total,
        "total_categories": n_class,
        "has_smiles": smi,
        "class_distribution": class_dist,
        "cross_database": cross,
    }

@app.get("/api/natural/filters")
def natural_filters():
    conn = get_db()
    classes = [r[0] for r in conn.execute(
        "SELECT DISTINCT std_class FROM natural_compounds ORDER BY std_class")]
    plants = set()
    for r in conn.execute("SELECT representative_plants FROM natural_compounds WHERE representative_plants IS NOT NULL"):
        for p in str(r[0]).replace("；",";").replace(",",";").split(";"):
            p = p.strip()
            if p:
                plants.add(p)
    conn.close()
    return {"classes": classes, "plants": sorted(plants)}

@app.get("/api/natural/compounds")
def natural_compounds(keyword: str = "", std_class: str = "", cas: str = "",
                      plants: str = "", has_smiles: bool = False,
                      page: int = 1, page_size: int = 50):
    conn = get_db()
    where, params = [], []
    if keyword:
        kw = f"%{keyword}%"
        where.append("(name_en LIKE ? OR name_cn LIKE ? OR cas LIKE ? OR formula LIKE ?)")
        params += [kw, kw, kw, kw]
    if std_class:
        where.append("std_class=?"); params.append(std_class)
    if cas:
        where.append("cas LIKE ?"); params.append(f"%{cas}%")
    if has_smiles:
        where.append("(smiles IS NOT NULL OR foodb_smiles IS NOT NULL)")
    if plants:
        pl = [p.strip() for p in plants.split(",") if p.strip()]
        if pl:
            ors = ["representative_plants LIKE ?" for _ in pl]
            params += [f"%{p}%" for p in pl]
            where.append("(" + " OR ".join(ors) + ")")
    ws = ("WHERE " + " AND ".join(where)) if where else ""
    total = conn.execute(f"SELECT COUNT(*) FROM natural_compounds {ws}", params).fetchone()[0]
    off = (page - 1) * page_size
    rows = conn.execute(f"""
        SELECT name_en,name_cn,std_class,raw_class,secondary_class,cas,formula,
               mce,targetmol,phenol_explorer,foodb,
               representative_plants,common_sources
        FROM natural_compounds {ws}
        ORDER BY name_en LIMIT ? OFFSET ?""", params + [page_size, off]).fetchall()
    conn.close()
    return {"total": total, "page": page, "page_size": page_size,
            "items": [dict(r) for r in rows]}

@app.get("/api/natural/compounds/{name}")
def natural_detail(name: str):
    conn = get_db()
    r = conn.execute("SELECT * FROM natural_compounds WHERE name_en=?", [name]).fetchone()
    if not r:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    data = {k: r[k] for k in r.keys()}
    conn.close()
    return data


@app.get("/api/oil/stats")
def oil_stats():
    conn = get_db()
    c = conn.cursor()
    total = c.execute("SELECT COUNT(*) FROM oil_compounds").fetchone()[0]
    n_class = c.execute("SELECT COUNT(DISTINCT std_class) FROM oil_compounds").fetchone()[0]
    ev = c.execute("SELECT COUNT(*) FROM oil_evidence").fetchone()[0]
    up = c.execute("SELECT COUNT(*) FROM oil_uplc").fetchone()[0]
    lit = c.execute("SELECT COUNT(*) FROM oil_literature").fetchone()[0]
    class_dist = {}
    for r in c.execute("SELECT std_class, COUNT(*) n FROM oil_compounds GROUP BY std_class ORDER BY n DESC"):
        class_dist[r[0]] = r[1]
    plant_cnt = {}
    for r in c.execute("SELECT sources FROM oil_compounds WHERE sources IS NOT NULL"):
        for p in str(r[0]).replace("；", ";").split(";"):
            p = p.strip()
            if p:
                plant_cnt[p] = plant_cnt.get(p, 0) + 1
    top_plants = sorted(plant_cnt.items(), key=lambda x: -x[1])[:20]
    conn.close()
    return {
        "total_compounds": total,
        "total_categories": n_class,
        "evidence_count": ev,
        "uplc_count": up,
        "literature_count": lit,
        "class_distribution": class_dist,
        "top_plants": [{"name": k, "count": v} for k, v in top_plants],
    }

@app.get("/api/oil/filters")
def oil_filters():
    conn = get_db()
    classes = [r[0] for r in conn.execute(
        "SELECT DISTINCT std_class FROM oil_compounds ORDER BY std_class")]
    plants = set()
    for r in conn.execute("SELECT source FROM oil_evidence WHERE source IS NOT NULL"):
        for p in str(r[0]).replace("；", ";").split(";"):
            p = p.strip()
            if p:
                plants.add(p)
    conn.close()
    return {"classes": classes, "plants": sorted(plants)}

@app.get("/api/oil/compounds")
def oil_compounds(keyword: str = "", std_class: str = "", plants: str = "",
                  has_evidence: bool = False, has_uplc: bool = False,
                  page: int = 1, page_size: int = 50):
    conn = get_db()
    where, params = [], []
    if keyword:
        kw = f"%{keyword}%"
        where.append("(name LIKE ? OR reference LIKE ?)")
        params += [kw, kw]
    if std_class:
        where.append("std_class=?"); params.append(std_class)
    if has_evidence:
        where.append("total_count > 0")
    if has_uplc:
        where.append("uplc_count > 0")
    if plants:
        pl = [p.strip() for p in plants.split(",") if p.strip()]
        if pl:
            ors = ["name IN (SELECT compound_name FROM oil_evidence WHERE source LIKE ?)" for _ in pl]
            params += [f"%{p}%" for p in pl]
            where.append("(" + " OR ".join(ors) + ")")
    ws = ("WHERE " + " AND ".join(where)) if where else ""
    total = conn.execute(f"SELECT COUNT(*) FROM oil_compounds {ws}", params).fetchone()[0]
    off = (page - 1) * page_size
    rows = conn.execute(f"""
        SELECT * FROM oil_compounds {ws}
        ORDER BY total_count DESC, name LIMIT ? OFFSET ?""",
        params + [page_size, off]).fetchall()
    conn.close()
    return {"total": total, "page": page, "page_size": page_size,
            "items": [dict(r) for r in rows]}

@app.get("/api/oil/compounds/{name}")
def oil_detail(name: str):
    conn = get_db()
    r = conn.execute("SELECT * FROM oil_compounds WHERE name=?", [name]).fetchone()
    if not r:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    data = {k: r[k] for k in r.keys()}
    data["evidence_list"] = [dict(x) for x in conn.execute(
        "SELECT * FROM oil_evidence WHERE compound_name=? ORDER BY id", [name])]
    data["uplc_list"] = [dict(x) for x in conn.execute(
        "SELECT * FROM oil_uplc WHERE compound_name=? ORDER BY id", [name])]
    conn.close()
    return data


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    content = await file.read()
    fname = file.filename or ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sns = wb.sheetnames
        if "Master_complete_only" in sns or "Master_strict_complete" in sns:
            added = parse_natural(wb)
        elif "合并汇总" in sns or "补充明细" in sns:
            added = parse_oil(wb)
        else:
            wb.close()
            return JSONResponse({"error": "无法识别文件格式"}, status_code=400)
        wb.close()
    except Exception as e:
        import traceback
        return JSONResponse({"error": str(e), "detail": traceback.format_exc()}, status_code=400)
    return {"status": "ok", "filename": fname, "added": added}


def parse_natural(wb):
    conn = sqlite3.connect(DB_PATH); c = conn.cursor(); n = 0
    for sn in ("Master_complete_only", "Master_strict_complete"):
        if sn not in wb.sheetnames: continue
        ws = wb[sn]; it = ws.iter_rows(values_only=True)
        header = next(it); col = {str(h).strip(): i for i, h in enumerate(header) if h}
        for row in it:
            name = _safe(row[col["name_en"]]) if "name_en" in col else None
            if not name: continue
            raw_cls = _safe(row[col["primary_class"]]) if "primary_class" in col else None
            std = norm_class(raw_cls, NATURAL_CLASS_MAP)
            c.execute("""INSERT OR REPLACE INTO natural_compounds
                (name_en,name_cn,std_class,raw_class,secondary_class,family,inclusion_level,
                 cas,formula,representative_plants,common_sources,mce,targetmol,
                 phenol_explorer,foodb,mce_url,targetmol_url,pe_url,smiles,smiles_source,
                 foodb_smiles,foodb_public_id,foodb_evidence_level,foodb_superklass,
                 foodb_klass,foodb_subklass,foodb_description,pubchem_title)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                [name,
                 _safe(row[col["name_cn"]]) if "name_cn" in col else None,
                 std, raw_cls,
                 _safe(row[col["secondary_class"]]) if "secondary_class" in col else None,
                 _safe(row[col["family"]]) if "family" in col else None,
                 _safe(row[col["inclusion_level"]]) if "inclusion_level" in col else None,
                 _safe(row[col["cas"]]) if "cas" in col else None,
                 _safe(row[col["formula"]]) if "formula" in col else None,
                 _safe(row[col["representative_plants"]]) if "representative_plants" in col else None,
                 _safe(row[col["common_sources"]]) if "common_sources" in col else None,
                 _safe(row[col["mce_phenols"]]) if "mce_phenols" in col else None,
                 _safe(row[col["targetmol_polyphenols"]]) if "targetmol_polyphenols" in col else None,
                 _safe(row[col["phenol_explorer"]]) if "phenol_explorer" in col else None,
                 _safe(row[col["foodb_phenolics"]]) if "foodb_phenolics" in col else None,
                 _safe(row[col["mce_url"]]) if "mce_url" in col else None,
                 _safe(row[col["targetmol_url"]]) if "targetmol_url" in col else None,
                 _safe(row[col["phenol_explorer_url"]]) if "phenol_explorer_url" in col else None,
                 _safe(row[col["smiles"]]) if "smiles" in col else None,
                 _safe(row[col["smiles_source"]]) if "smiles_source" in col else None,
                 _safe(row[col["foodb_smiles"]]) if "foodb_smiles" in col else None,
                 _safe(row[col["foodb_public_id"]]) if "foodb_public_id" in col else None,
                 _safe(row[col["foodb_evidence_level"]]) if "foodb_evidence_level" in col else None,
                 _safe(row[col["foodb_superklass"]]) if "foodb_superklass" in col else None,
                 _safe(row[col["foodb_klass"]]) if "foodb_klass" in col else None,
                 _safe(row[col["foodb_subklass"]]) if "foodb_subklass" in col else None,
                 _safe(row[col["foodb_description"]]) if "foodb_description" in col else None,
                 _safe(row[col["pubchem_title"]]) if "pubchem_title" in col else None])
            n += 1
        if sn == "Master_complete_only": break
    conn.commit(); conn.close()
    return {"natural_compounds": n}


def parse_oil(wb):
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    counts = {"oil_compounds": 0, "oil_evidence": 0, "oil_uplc": 0, "oil_literature": 0}

    if "合并汇总" in wb.sheetnames:
        ws = wb["合并汇总"]; it = ws.iter_rows(values_only=True)
        header = next(it); col = {str(h).strip(): i for i, h in enumerate(header) if h}
        for row in it:
            name = _safe(row[col["化合物名称"]]) if "化合物名称" in col else None
            if not name: continue
            raw_cls = _safe(row[col["类别"]]) if "类别" in col else None
            std = norm_class(raw_cls, OIL_CLASS_MAP)
            c.execute("""INSERT OR REPLACE INTO oil_compounds
                (name,std_class,raw_class,sources,reference,total_count,orig_count,
                 new_lit_count,uplc_count,online_lit_count,online_entity_count)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                [name, std, raw_cls,
                 _safe(row[col["来源"]]) if "来源" in col else None,
                 _safe(row[col["参考文献"]]) if "参考文献" in col else None,
                 _int(row[col["出现次数"]]) if "出现次数" in col else 0,
                 _int(row[col["原表出现次数"]]) if "原表出现次数" in col else 0,
                 _int(row[col["新增文献次数"]]) if "新增文献次数" in col else 0,
                 _int(row[col["UPLC出现次数"]]) if "UPLC出现次数" in col else 0,
                 _int(row[col["在线文献次数"]]) if "在线文献次数" in col else 0,
                 _int(row[col["在线化学实体次数"]]) if "在线化学实体次数" in col else 0])
            counts["oil_compounds"] += 1

    if "补充明细" in wb.sheetnames:
        ws = wb["补充明细"]; it = ws.iter_rows(values_only=True)
        header = next(it); col = {str(h).strip(): i for i, h in enumerate(header) if h}
        for row in it:
            name = _safe(row[col["化合物名称"]]) if "化合物名称" in col else None
            if not name: continue
            raw_cls = _safe(row[col["类别"]]) if "类别" in col else None
            std = norm_class(raw_cls, OIL_CLASS_MAP)
            c.execute("""INSERT INTO oil_evidence
                (compound_name,std_class,raw_class,source,reference,data_source,
                 platform,pdf_file,doi,url,snippet)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                [name, std, raw_cls,
                 _safe(row[col["来源"]]) if "来源" in col else None,
                 _safe(row[col["参考文献"]]) if "参考文献" in col else None,
                 _safe(row[col["数据来源"]]) if "数据来源" in col else None,
                 _safe(row[col["来源平台"]]) if "来源平台" in col else None,
                 _safe(row[col["PDF文件"]]) if "PDF文件" in col else None,
                 _safe(row[col["DOI"]]) if "DOI" in col else None,
                 _safe(row[col["文献链接"]]) if "文献链接" in col else None,
                 _safe(row[col["证据片段"]]) if "证据片段" in col else None])
            counts["oil_evidence"] += 1

    if "菜籽UPLC筛选" in wb.sheetnames:
        ws = wb["菜籽UPLC筛选"]; it = ws.iter_rows(values_only=True)
        header = next(it); col = {str(h).strip(): i for i, h in enumerate(header) if h}
        for row in it:
            name = _safe(row[col["化合物名称"]]) if "化合物名称" in col else None
            if not name: continue
            raw_cls = _safe(row[col["类别"]]) if "类别" in col else None
            std = norm_class(raw_cls, OIL_CLASS_MAP)
            c.execute("""INSERT INTO oil_uplc
                (compound_name,std_class,original_name,raw_class,primary_class,
                 secondary_class,formula,q1,q3,mw,ionization,level,cas,cpd_id,
                 fp,ep,kegg_map)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                [name, std,
                 _safe(row[col["原始化合物名称"]]) if "原始化合物名称" in col else None,
                 raw_cls,
                 _safe(row[col["一级分类"]]) if "一级分类" in col else None,
                 _safe(row[col["二级分类"]]) if "二级分类" in col else None,
                 _safe(row[col["Formula"]]) if "Formula" in col else None,
                 _safe(row[col["Q1 (Da)"]]) if "Q1 (Da)" in col else None,
                 _safe(row[col["Q3 (Da)"]]) if "Q3 (Da)" in col else None,
                 _safe(row[col["Molecular weight (Da)"]]) if "Molecular weight (Da)" in col else None,
                 _safe(row[col["Ionization model"]]) if "Ionization model" in col else None,
                 _safe(row[col["Level"]]) if "Level" in col else None,
                 _safe(row[col["CAS"]]) if "CAS" in col else None,
                 _safe(row[col["cpd_ID"]]) if "cpd_ID" in col else None,
                 _safe(row[col["FP"]]) if "FP" in col else None,
                 _safe(row[col["EP"]]) if "EP" in col else None,
                 _safe(row[col["kegg_map"]]) if "kegg_map" in col else None])
            counts["oil_uplc"] += 1

    if "在线化学实体记录" in wb.sheetnames:
        ws = wb["在线化学实体记录"]; it = ws.iter_rows(values_only=True)
        header = next(it); col = {str(h).strip(): i for i, h in enumerate(header) if h}
        for row in it:
            title = _safe(row[col["题名"]]) if "题名" in col else None
            if not title: continue
            c.execute("""INSERT INTO oil_literature
                (crop,platform,query_text,title,doi,year,journal,url,
                 compounds_identified,count)
                VALUES (?,?,?,?,?,?,?,?,?,?)""",
                [_safe(row[col["检索作物"]]) if "检索作物" in col else None,
                 _safe(row[col["来源平台"]]) if "来源平台" in col else None,
                 _safe(row[col["检索词"]]) if "检索词" in col else None,
                 title,
                 _safe(row[col["DOI"]]) if "DOI" in col else None,
                 _safe(row[col["年份"]]) if "年份" in col else None,
                 _safe(row[col["期刊"]]) if "期刊" in col else None,
                 _safe(row[col["文献链接"]]) if "文献链接" in col else None,
                 _safe(row[col["识别化合物"]]) if "识别化合物" in col else None,
                 _safe(row[col["识别数量"]]) if "识别数量" in col else None])
            counts["oil_literature"] += 1

    conn.commit(); conn.close()
    return counts


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.environ.get("PORT", 8000)))
